import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
os.chdir(os.path.dirname(os.path.abspath(__file__)))
import consultas
E = json.load(open('04_incluidos.json', encoding='utf8'))
d = json.load(open('03_resultado.json', encoding='utf8'))
F, O, cif = d['final'], d['otros'], d['cifras']
base = json.load(open('02_para_cribado.json', encoding='utf8'))
R = {r['clave']: r for r in base['registros']}
crib = {}
import glob
for f in glob.glob('decisiones/lote_*.jsonl'):
    for l in open(f, encoding='utf8'):
        if l.strip():
            x = json.loads(l); crib[x['clave']] = x

TEMA = {'defensa': 'Defensa', 'ofensiva': 'Ofensiva', 'educacion': 'Educación', 'seguridad_agentes': 'Seguridad de los agentes', 'general': 'General'}
TIPO = {'experimental': 'Experimental', 'propuesta_sin_evaluacion': 'Propuesta sin evaluación', 'revision_sistematica': 'Revisión sistemática',
        'revision_narrativa': 'Revisión narrativa', 'posicion': 'Artículo de posición'}
TEC = {'llm_agente': 'Agente basado en LLM', 'multiagente_llm': 'Sistema multiagente con LLM', 'rl': 'Aprendizaje por refuerzo', 'mixta': 'Mixta',
       'llm_sin_agencia': 'IA generativa sin agencia'}
CRIT = {'IA': 'Incluido, agentes de IA en ciberseguridad', 'IB': 'Incluido, IA en educación en ciberseguridad', 'E1': 'No trata ciberseguridad',
        'E2': 'Agente no se refiere a IA', 'E3': 'IA no agéntica', 'E4': 'Educación sin IA', 'E5': 'No es trabajo académico',
        'E6': 'No cumple los criterios en texto completo', 'E7': 'No es un trabajo completo', 'E8': 'Versión duplicada', 'NR': 'Texto completo no recuperado',
        'DUDA': 'Incierto en cribado'}

wb = Workbook()
cab = Font(bold=True, color='FFFFFF'); relleno = PatternFill('solid', fgColor='1F3864')
def hoja(ws, filas, anchos):
    for c in ws[1]:
        c.font = cab; c.fill = relleno; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

ws = wb.active; ws.title = 'Incluidos'
ws.append(['N.º', 'Clave', 'Autores', 'Año', 'Título', 'Fuente', 'Revisión por pares', 'Origen', 'Tema', 'Tipo de estudio', 'Tecnología',
           'Población (educación)', 'Hallazgo principal', 'DOI', 'Enlace'])
for i, e in enumerate(E, 1):
    ws.append([i, e['clave'], e['autores'], e['anio'], e['titulo'], e['venue'], 'Sí' if e['revisado'] else 'No (preprint)', e['origen'],
               TEMA.get(e['tema'], e['tema']), TIPO.get(e['tipo_estudio'], e['tipo_estudio']), TEC.get(e['tecnologia'], e['tecnologia']),
               e['poblacion'] if e['tema'] == 'educacion' else '', e['hallazgo'], e['doi'], e['url']])
hoja(ws, None, [6, 8, 22, 6, 60, 30, 14, 14, 20, 22, 26, 16, 60, 26, 40])

ws = wb.create_sheet('Registros cribados')
ws.append(['Clave', 'Título', 'Año', 'Tipo', 'Fuente', 'Cadena(s)', 'Decisión en cribado', 'Motivo en cribado', 'Nota del cribado',
           'Decisión final', 'Motivo final', 'Texto completo leído'])
for k in sorted(F):
    r, v, c = R[k], F[k], crib[k]
    ws.append([k, r['titulo'], r['anio'], r['tipo'], r['venue'], ', '.join(r['cadenas']), c['decision'], CRIT.get(c['criterio'], c['criterio']),
               c.get('nota', ''), v['decision'], CRIT.get(v['criterio'], v['criterio']), v.get('fuente_leida', '')])
hoja(ws, None, [8, 60, 6, 14, 30, 10, 14, 30, 40, 14, 34, 40])

ws = wb.create_sheet('Otros métodos')
ws.append(['Clave', 'Autores', 'Año', 'Título', 'Fuente', 'Decisión', 'Motivo', 'Hallazgo o motivo'])
for k, v in O.items():
    ws.append([k, v['autores'], v['anio'], v['titulo'], v['venue'], v['decision'], CRIT.get(v['criterio'], v['criterio']), v['hallazgo']])
hoja(ws, None, [8, 24, 6, 60, 30, 10, 34, 60])

ws = wb.create_sheet('Protocolo')
ws.append(['Elemento', 'Descripción'])
filas = [
    ('Pregunta de la revisión', '¿Qué aplicaciones, riesgos y experiencias educativas se han reportado sobre los agentes de IA en ciberseguridad entre 2022 y 2026?'),
    ('Fecha de la búsqueda', base['fecha_busqueda']),
    ('Fuentes', 'OpenAlex (índice abierto con registros de IEEE, ACM, Springer, Elsevier, MDPI y otras editoriales) y arXiv. Semantic Scholar se usó para completar resúmenes y localizar textos completos.'),
    ('Periodo', '1 de enero de 2022 a 30 de septiembre de 2026'),
    ('Cadena A en OpenAlex (agentes)', consultas.OPENALEX['A']),
    ('Cadena B en OpenAlex (educación)', consultas.OPENALEX['B']),
    ('Cadena A en arXiv', consultas.ARXIV['A']),
    ('Cadena B en arXiv', consultas.ARXIV['B']),
    ('Identificados', f"OpenAlex {cif['identificados']['OpenAlex']}, arXiv {cif['identificados']['arXiv']}"),
    ('Eliminados antes del cribado', f"Duplicados {cif['duplicados']}, tipo no académico {cif['fuera_tipo']}, idioma {cif['fuera_idioma']}"),
    ('Cribados', cif['cribados']), ('Excluidos en cribado', cif['excluidos_cribado']),
    ('Buscados', cif['buscados']), ('No recuperados', f"{cif['no_recuperados']} (sin acceso abierto {cif['no_recuperados_sin_oa']}, acceso abierto no descargable {cif['no_recuperados_oa_bloqueado']})"),
    ('Evaluados en texto completo', cif['evaluados']), ('Incluidos de bases de datos', cif['incluidos_bd']),
    ('Otros métodos', f"Identificados {cif['otros_identificados']}, ya en bases {cif['otros_ya_en_bd']}, evaluados {cif['otros_evaluados']}, incluidos {cif['otros_incluidos']}"),
    ('Total incluidos', cif['incluidos_total']),
    ('Procedimiento de cribado', 'Cribado por título y resumen y evaluación de texto completo realizados con apoyo de modelos de IA bajo criterios escritos (IA, IB, E1 a E8), con auditoría de una muestra aleatoria de decisiones.'),
]
for f in filas: ws.append([f[0], str(f[1])])
hoja(ws, None, [34, 140])
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
destino = r'C:\Users\SISTEMAS10\Desktop\Estudio\TESIS NUEVA\Revision sistematica PRISMA - datos.xlsx'
wb.save(destino)
print('guardado', destino, len(E), len(F))
